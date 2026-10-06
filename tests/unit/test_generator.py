import csv
import hashlib
from collections import Counter, defaultdict
from decimal import Decimal

from openpyxl import load_workbook

from erp_migration.config import Paths
from erp_migration.generate import generate
from erp_migration.generate.catalog import make_catalog
from erp_migration.generate.simulate import simulate


def _digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_same_seed_gives_the_same_books_and_truth(tmp_path):
    a, b = Paths(tmp_path / "a"), Paths(tmp_path / "b")
    generate(a, 7, 2018, 2019)
    generate(b, 7, 2018, 2019)
    for name in ("rows.csv", "supplier_aliases.csv", "suppliers.csv", "manifest.json"):
        assert _digest(a.truth / name) == _digest(b.truth / name)
    generate(b, 8, 2018, 2019)
    assert _digest(a.truth / "rows.csv") != _digest(b.truth / "rows.csv")


def test_books_contain_the_promised_mess(small_dataset):
    paths = small_dataset
    rows = list(csv.DictReader((paths.truth / "rows.csv").open()))
    kinds = Counter(r["kind"] for r in rows)
    reasons = Counter(r["expected_reason"] for r in rows if r["expected_outcome"] == "rejected")
    assert kinds["blank"] and kinds["subtotal"] and kinds["data"]
    assert {"BAD_DATE", "BAD_AMOUNT", "MISSING_FIELD", "UNKNOWN_PRODUCT"} <= set(reasons)
    assert any(r["expected_outcome"] == "merged" for r in rows)

    date_cells: Counter[str] = Counter()
    merged_ranges = 0
    for book in paths.books.glob("purchases_*.xlsx"):
        wb = load_workbook(book)
        for ws in wb.worksheets:
            merged_ranges += len(ws.merged_cells.ranges)
            for (value,) in ws.iter_rows(min_col=1, max_col=1, values_only=True):
                date_cells[type(value).__name__] += 1
    assert merged_ranges > 0
    assert date_cells["datetime"] and date_cells["str"]


def test_every_spelling_belongs_to_one_supplier(small_dataset):
    aliases = list(csv.DictReader((small_dataset.truth / "supplier_aliases.csv").open()))
    assert len({a["alias"] for a in aliases}) == len(aliases)
    by_supplier = defaultdict(set)
    for a in aliases:
        by_supplier[a["supplier_code"]].add(a["alias"])
    assert max(len(v) for v in by_supplier.values()) > 1  # suppliers really are misspelt


def test_true_stock_never_goes_negative():
    catalog = make_catalog(42, 2015, 2017)
    ledger = simulate(catalog, 42, 2015, 2017)
    daily: dict[str, Counter] = defaultdict(Counter)
    for m in ledger.movements:
        daily[m.sku][m.movement_date] += m.quantity
    for days in daily.values():
        balance = Decimal(0)
        for d in sorted(days):
            balance += days[d]
            assert balance >= 0


def test_payroll_is_internally_consistent():
    catalog = make_catalog(42, 2015, 2016)
    ledger = simulate(catalog, 42, 2015, 2016)
    assert all(p.net == p.gross - p.deductions for p in ledger.payroll)
    assert len({(p.period, p.employee_code) for p in ledger.payroll}) == len(ledger.payroll)
