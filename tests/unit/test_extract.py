from datetime import datetime

from openpyxl import Workbook

from erp_migration.etl.extract import domain_of, extract_book, normalize_label, year_of


def _two_level_book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Q1"
    ws["A1"] = "PURCHASES BOOK 2019"
    ws.merge_cells("A1:G1")
    ws["A2"], ws["C2"], ws["D2"], ws["F2"] = "Order", "Supplier", "Item", "Amount (Frw)"
    ws.merge_cells("A2:B2")
    ws.merge_cells("C2:C3")
    ws.merge_cells("D2:E2")
    ws.merge_cells("F2:G2")
    for col, label in zip(
        "ABDEFG", ["Date", "Number", "Code", "Qty", "Value", "Currency"], strict=True
    ):
        ws[f"{col}3"] = label
    rows = [
        [datetime(2019, 1, 4), "PO-2019-00001", "Walker Textiles", "FAB-COT-001", 10, 5000, "RWF"],
        [None] * 7,
        ["Total January 2019", None, None, None, None, 5000, None],
        ["Date", "Number", "Supplier", "Code", "Qty", "Value", "Currency"],
        ["05/02/2019", "PO-2019-00002", "Total Packaging Ltd", "PKG-BAG-001", 5, "2,500/=", None],
    ]
    for i, values in enumerate(rows, start=4):
        for j, v in enumerate(values, start=1):
            if v is not None:
                ws.cell(i, j, v)
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Figures checked against supplier statements."
    wb.save(path)


def test_two_level_merged_header_is_read(tmp_path):
    path = tmp_path / "purchases_2019.xlsx"
    _two_level_book(path)
    q1, notes = extract_book(path)
    assert q1.layout.header_row == 3
    assert sorted(q1.layout.columns.values()) == sorted(
        ["date", "po_number", "supplier", "sku", "quantity", "amount", "currency"]
    )
    assert q1.layout.hints == {"amount": "RWF", "currency": "RWF"}
    assert notes.skipped_reason == "no recognisable header"


def test_rows_are_classified(tmp_path):
    path = tmp_path / "purchases_2019.xlsx"
    _two_level_book(path)
    q1, _ = extract_book(path)
    kinds = [(r.row_num, r.kind) for r in q1.rows]
    assert kinds == [
        (4, "data"),
        (5, "blank"),
        (6, "subtotal"),
        (7, "repeated_header"),
        (8, "data"),
    ]
    # A supplier called "Total ..." is data, because the row has a real PO number.
    assert q1.rows[-1].cells["supplier"] == "Total Packaging Ltd"
    assert q1.rows[0].cells["date"] == datetime(2019, 1, 4)


def test_file_names_and_labels():
    assert domain_of("purchases_2019_revised.xlsx") == "purchases"
    assert domain_of("item_master.xlsx") == "products"
    assert domain_of("budget_2019.xlsx") is None
    assert year_of("purchases_2019_revised.xlsx") == 2019
    assert year_of("item_master.xlsx") is None
    assert normalize_label("Amount (Frw)") == ("amount", "RWF")
    assert normalize_label("Order No.") == ("order no", None)
    assert normalize_label("Unit price (US$)") == ("unit price", "USD")
