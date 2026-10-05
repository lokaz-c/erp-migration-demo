"""Write assembled books to .xlsx files and derive the ground truth from them."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from erp_migration.generate.books import Book, Sheet
from erp_migration.generate.render import Row

NOISE_REASONS = {
    "blank": "BLANK_ROW",
    "subtotal": "SUBTOTAL_ROW",
    "repeated_header": "REPEATED_HEADER",
}


@dataclass
class Placement:
    book: str
    sheet_index: int
    sheet: str
    row_num: int
    domain: str
    row: Row


def _write_header(ws, sheet: Sheet) -> int:
    """Write title and header rows; return the first data row number."""
    n = len(sheet.columns)
    r = 1
    if sheet.layout in ("title", "two_level"):
        ws.cell(1, 1, sheet.title).font = Font(bold=True)
        if n > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        r = 2
        if sheet.layout == "title":
            ws.cell(2, 1, "Kept by: accounts office")
            r = 4
    bold = Font(bold=True)
    if sheet.layout == "two_level" and any(c.parent for c in sheet.columns):
        c = 1
        while c <= n:
            col = sheet.columns[c - 1]
            if col.parent is None:
                ws.cell(r, c, col.label).font = bold
                ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
                c += 1
                continue
            end = c
            while end < n and sheet.columns[end].parent == col.parent:
                end += 1
            ws.cell(r, c, col.parent).font = bold
            if end > c:
                ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=end)
            for k in range(c, end + 1):
                ws.cell(r + 1, k, sheet.columns[k - 1].label).font = bold
            c = end + 1
        return r + 2
    for c, col in enumerate(sheet.columns, start=1):
        ws.cell(r, c, col.label).font = bold
    return r + 1


def write_book(book: Book, out_dir: Path) -> list[Placement]:
    wb = Workbook()
    wb.remove(wb.active)
    placements: list[Placement] = []
    for index, sheet in enumerate(book.sheets):
        ws = wb.create_sheet(sheet.name)
        if sheet.free_text is not None:
            for line in sheet.free_text:
                ws.append(line)
            continue
        first = _write_header(ws, sheet)
        for i, row in enumerate(sheet.rows):
            excel_row = first + i
            for c, col in enumerate(sheet.columns, start=1):
                value = row.values.get(col.field)
                if value is None:
                    continue
                cell = ws.cell(excel_row, c, value)
                fmt = row.formats.get(col.field)
                if fmt:
                    cell.number_format = fmt
            placements.append(
                Placement(book.filename, index, sheet.name, excel_row, book.domain, row)
            )
        for c in range(1, len(sheet.columns) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16
    wb.properties.creator = "erp-migration-demo synthetic generator"
    wb.save(out_dir / book.filename)
    return placements


def expected_outcomes(placements: list[Placement]) -> dict[int, tuple[str, str]]:
    """What a careful human would do with each row, given the true records.

    The first valid occurrence of a record (by book, sheet, row) is loaded,
    later identical copies are merged, and copies whose values disagree with
    the first are rejected as conflicting duplicates.
    """
    outcome: dict[int, tuple[str, str]] = {}
    by_record: dict[str, list[Placement]] = defaultdict(list)
    for p in placements:
        if p.row.kind in NOISE_REASONS:
            outcome[id(p)] = ("skipped", NOISE_REASONS[p.row.kind])
        elif p.row.reason:
            outcome[id(p)] = ("rejected", p.row.reason)
        else:
            by_record[p.row.record_id].append(p)
    for ps in by_record.values():
        ps.sort(key=lambda p: (p.book, p.sheet_index, p.row_num))
        first = ps[0]
        outcome[id(first)] = ("loaded", "")
        for p in ps[1:]:
            same = p.row.signature == first.row.signature
            outcome[id(p)] = ("merged", "") if same else ("rejected", "CONFLICTING_DUPLICATE")
    return outcome


TRUTH_COLUMNS = [
    "book",
    "sheet_index",
    "sheet",
    "row_num",
    "domain",
    "kind",
    "record_id",
    "expected_outcome",
    "expected_reason",
    "true_date",
    "true_currency",
    "true_amount",
    "supplier_code",
    "supplier_raw",
]


def write_truth(
    placements: list[Placement], suppliers, truth_dir: Path, manifest: dict
) -> dict[str, int]:
    truth_dir.mkdir(parents=True, exist_ok=True)
    outcome = expected_outcomes(placements)
    aliases: dict[str, str] = {}
    with (truth_dir / "rows.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(TRUTH_COLUMNS)
        for p in placements:
            o, reason = outcome[id(p)]
            t = p.row.truth
            w.writerow(
                [
                    p.book,
                    p.sheet_index,
                    p.sheet,
                    p.row_num,
                    p.domain,
                    p.row.kind,
                    p.row.record_id or "",
                    o,
                    reason,
                    t.get("true_date", ""),
                    t.get("true_currency", ""),
                    t.get("true_amount", ""),
                    t.get("supplier_code", ""),
                    t.get("supplier_raw", ""),
                ]
            )
            raw = t.get("supplier_raw")
            if (
                p.domain == "purchases"
                and raw
                and aliases.setdefault(raw, t["supplier_code"]) != t["supplier_code"]
            ):
                raise ValueError(f"spelling {raw!r} used for two suppliers")
    with (truth_dir / "supplier_aliases.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["alias", "supplier_code"])
        for alias in sorted(aliases):
            w.writerow([alias, aliases[alias]])
    with (truth_dir / "suppliers.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["supplier_code", "canonical_name", "category", "currency"])
        for s in suppliers:
            w.writerow([s.code, s.name, s.category, s.currency])
    counts: dict[str, int] = defaultdict(int)
    for o, _ in outcome.values():
        counts[o] += 1
    manifest = {
        **manifest,
        "rows": len(placements),
        "expected": dict(sorted(counts.items())),
        "supplier_aliases": len(aliases),
    }
    (truth_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return dict(counts)
