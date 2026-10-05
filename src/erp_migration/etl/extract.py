"""Read workbooks into rows of named fields.

Books were kept by different people, so nothing about the layout is assumed:

* the domain comes from the file name (purchases_2019.xlsx, payroll_2018.xlsx);
* the header row is found by scoring the first rows against a synonym list,
  so title rows, "Kept by" lines and blank rows above it are skipped;
* merged cells are expanded first, so a two-level header such as
  "Amount" over "Value | Currency" is read as "amount value" and
  "amount currency";
* a currency written in a header ("Amount (Frw)") is kept as a hint;
* sheets with no recognisable header (notes, summaries) are skipped and
  reported, not guessed at.

Every row below the header is kept and classified as data, blank, subtotal
or a repeated header, so the report can account for all of them.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from erp_migration.parsing.amounts import CURRENCY_WORDS
from erp_migration.parsing.text import clean, normalize_doc, normalize_employee_code, normalize_sku

HEADER_SEARCH_ROWS = 12

SYNONYMS: dict[str, dict[str, list[str]]] = {
    "purchases": {
        "date": [
            "date",
            "invoice date",
            "inv date",
            "date of purchase",
            "purchase date",
            "order date",
        ],
        "po_number": ["po no", "po number", "po #", "po", "order no", "order number", "order #"],
        "supplier": ["supplier", "supplier name", "vendor", "vendor name"],
        "sku": ["item code", "code", "product code", "sku"],
        "description": ["description", "item description", "item", "product"],
        "quantity": ["qty", "quantity"],
        "unit_price": ["unit price", "rate", "u/price", "price", "unit cost"],
        "amount": ["amount", "total", "value", "amount value"],
        "currency": ["currency", "curr", "ccy", "amount currency"],
    },
    "sales": {
        "date": ["date", "sale date", "invoice date"],
        "receipt_no": ["receipt no", "receipt #", "receipt number", "invoice no", "sale number"],
        "channel": ["channel", "sales channel", "outlet"],
        "sku": ["item code", "code", "product code", "sku"],
        "description": ["description", "item description", "item", "product"],
        "quantity": ["qty", "quantity"],
        "unit_price": ["unit price", "price", "u/price", "rate"],
        "amount": ["amount", "total", "value", "amount value"],
        "currency": ["currency", "curr", "ccy", "amount currency"],
    },
    "inventory": {
        "date": ["date", "movement date", "document date"],
        "voucher_no": ["voucher no", "voucher", "doc no", "ref no", "document number"],
        "sku": ["item code", "code", "product code", "sku"],
        "description": ["description", "item", "product", "item description"],
        "movement_type": ["movement", "type", "movement type"],
        "qty_in": ["qty in", "received", "quantity in", "in"],
        "qty_out": ["qty out", "issued", "quantity out", "out"],
        "quantity": ["quantity", "qty"],
        "reference": ["reference", "ref", "remarks"],
    },
    "payroll": {
        "employee_code": ["emp no", "staff id", "employee id", "employee no", "id"],
        "name": ["employee name", "name", "names"],
        "position": ["position", "job title", "function", "employee position"],
        "gross": ["gross salary", "gross", "gross pay"],
        "deductions": ["deductions", "total deductions"],
        "net": ["net pay", "net", "net salary"],
    },
    "products": {
        "sku": ["item code", "code", "sku", "product code"],
        "description": ["description", "item description"],
        "category": ["category"],
        "uom": ["unit", "uom", "unit of measure"],
    },
}

# A sheet is only read if its header has these (any one field of each group).
REQUIRED: dict[str, list[set[str]]] = {
    "purchases": [{"date"}, {"po_number"}, {"supplier"}, {"sku", "description"}, {"amount"}],
    "sales": [{"date"}, {"receipt_no"}, {"sku", "description"}, {"amount"}],
    "inventory": [{"date"}, {"voucher_no"}, {"sku"}, {"movement_type"}, {"quantity", "qty_in"}],
    "payroll": [{"name"}, {"gross"}, {"net"}],
    "products": [{"sku"}, {"description"}],
}
KEY_FIELD = {
    "purchases": "po_number",
    "sales": "receipt_no",
    "inventory": "voucher_no",
    "payroll": "employee_code",
    "products": "sku",
}
_SUBTOTAL = re.compile(
    r"^\s*((grand\s+)?(sub[\s-]?)?total\b|balance\s+[bc]/f|(carried|brought)\s+forward)", re.I
)
_HINT = re.compile(r"\(\s*([^)]*)\s*\)")


def domain_of(filename: str) -> str | None:
    name = filename.lower()
    if name.startswith("item_master"):
        return "products"
    for domain in ("purchases", "sales", "inventory", "payroll"):
        if name.startswith(domain + "_"):
            return domain
    return None


def year_of(filename: str) -> int | None:
    m = re.search(r"_(\d{4})(?:[_.]|$)", filename)
    return int(m[1]) if m else None


def normalize_label(text: Any) -> tuple[str, str | None]:
    """Header text -> (lookup label, currency hint). 'Amount (Frw)' -> ('amount', 'RWF')."""
    raw = clean(text).lower()
    hint = None
    for inner in _HINT.findall(raw):
        hint = hint or CURRENCY_WORDS.get(inner.strip())
    label = _HINT.sub(" ", raw).replace(".", " ").replace(":", " ")
    return re.sub(r"\s+", " ", label).strip(), hint


@dataclass
class Layout:
    header_row: int
    columns: dict[int, str]  # 0-based column index -> field
    labels: dict[int, str]  # 0-based column index -> normalized child label
    hints: dict[str, str] = field(default_factory=dict)  # field -> currency from header


@dataclass
class RawRow:
    book: str
    sheet: str
    sheet_index: int
    row_num: int
    domain: str
    kind: str  # data | blank | subtotal | repeated_header
    cells: dict[str, Any]


@dataclass
class SheetData:
    book: str
    sheet: str
    sheet_index: int
    domain: str
    year: int | None
    layout: Layout | None
    rows: list[RawRow] = field(default_factory=list)
    skipped_reason: str | None = None


def _grid(ws) -> list[list[Any]]:
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max((len(r) for r in grid), default=0)
    for r in grid:
        r.extend([None] * (width - len(r)))
    for rng in ws.merged_cells.ranges:
        value = grid[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row - 1, rng.max_row):
            for c in range(rng.min_col - 1, rng.max_col):
                grid[r][c] = value
    return grid


def _map_row(domain: str, child: list[Any], parent: list[Any] | None) -> Layout:
    lookup = {syn: f for f, syns in SYNONYMS[domain].items() for syn in syns}
    columns: dict[int, str] = {}
    labels: dict[int, str] = {}
    hints: dict[str, str] = {}
    for i, cell in enumerate(child):
        label, hint = normalize_label(cell)
        if not label:
            continue
        p_label, p_hint = normalize_label(parent[i]) if parent else ("", None)
        candidates = [label]
        if p_label and p_label != label:
            candidates.insert(0, f"{p_label} {label}")
        fld = next((lookup[c] for c in candidates if c in lookup), None)
        if fld is None or fld in columns.values():
            continue
        columns[i] = fld
        labels[i] = label
        if hint or p_hint:
            hints[fld] = hint or p_hint
    return Layout(0, columns, labels, hints)


def find_layout(domain: str, grid: list[list[Any]]) -> Layout | None:
    best: Layout | None = None
    for r in range(min(HEADER_SEARCH_ROWS, len(grid))):
        layout = _map_row(domain, grid[r], grid[r - 1] if r > 0 else None)
        fields = set(layout.columns.values())
        if not all(group & fields for group in REQUIRED[domain]):
            continue
        if best is None or len(layout.columns) > len(best.columns):
            layout.header_row = r + 1
            best = layout
    return best


def classify(domain: str, cells: dict[str, Any], layout: Layout) -> str:
    texts = {f: clean(v) for f, v in cells.items()}
    if not any(texts.values()):
        return "blank"
    matches = sum(texts.get(f, "").lower() == layout.labels[i] for i, f in layout.columns.items())
    if matches >= max(2, int(0.6 * len(layout.columns))):
        return "repeated_header"
    if any(_SUBTOTAL.match(t) for t in texts.values() if t):
        key = cells.get(KEY_FIELD[domain])
        normalizer = {
            "purchases": lambda v: normalize_doc(v, "purchases"),
            "sales": lambda v: normalize_doc(v, "sales"),
            "inventory": lambda v: normalize_doc(v, "inventory"),
            "payroll": normalize_employee_code,
            "products": normalize_sku,
        }[domain]
        if not key or normalizer(key) is None:
            return "subtotal"
    return "data"


def extract_book(path: Path) -> list[SheetData]:
    domain = domain_of(path.name)
    if domain is None:
        return [
            SheetData(
                path.name,
                "*",
                0,
                "unknown",
                None,
                None,
                skipped_reason="file name does not identify a domain",
            )
        ]
    wb = load_workbook(path, data_only=True)
    out: list[SheetData] = []
    for index, ws in enumerate(wb.worksheets):
        grid = _grid(ws)
        layout = find_layout(domain, grid)
        sheet = SheetData(path.name, ws.title, index, domain, year_of(path.name), layout)
        if layout is None:
            sheet.skipped_reason = "no recognisable header"
            out.append(sheet)
            continue
        for r in range(layout.header_row, len(grid)):
            values = grid[r]
            cells = {f: values[i] for i, f in layout.columns.items() if i < len(values)}
            sheet.rows.append(
                RawRow(
                    path.name,
                    ws.title,
                    index,
                    r + 1,
                    domain,
                    classify(domain, cells, layout),
                    cells,
                )
            )
        out.append(sheet)
    wb.close()
    return out


def extract_all(books_dir: Path) -> list[SheetData]:
    sheets: list[SheetData] = []
    for path in sorted(books_dir.glob("*.xlsx")):
        sheets.extend(extract_book(path))
    return sheets
